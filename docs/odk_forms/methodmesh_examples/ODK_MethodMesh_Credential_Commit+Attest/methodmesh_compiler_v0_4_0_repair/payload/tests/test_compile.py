import json
from datetime import datetime, timezone, timedelta
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from methodmesh_xlsform.compiler import compile_xlsform
from methodmesh_xlsform.sentinel import run_form_compile_task
from methodmesh_xlsform.errors import ValidationError


class CompileTests(unittest.TestCase):
    def make_source(self, path: Path):
        wb = Workbook()
        ws = wb.active
        ws.title = "survey"
        ws.append(["type", "name", "label", "readonly", "mm_commit"])
        ws.append(["text", "participant_id", "Participant ID", None, None])
        ws.append(["integer", "age", "Age", "no", None])
        ws.append(["select_one outcome", "outcome", "Outcome", None, None])
        ws.append(["text", "ui_helper", "UI helper", None, "exclude"])
        ws.append(["calculate", "calc", None, None, None])
        ws["C6"] = None
        ws["A6"] = "calculate"
        ws["B6"] = "calc"
        # calculation column may be absent in source; compiler should add it.

        ch = wb.create_sheet("choices")
        ch.append(["list_name", "name", "label"])
        ch.append(["outcome", "normal", "Normal"])
        ch.append(["outcome", "abnormal", "Abnormal"])

        st = wb.create_sheet("settings")
        st.append(["form_title", "form_id", "version"])
        st.append(["Test", "test_form", "2026092201"])

        mm = wb.create_sheet("methodmesh")
        mm.append(["key", "value"])
        mm.append(["study_id", "TEST_STUDY"])
        mm.append(["timestamp_policy", "preferred"])
        wb.save(path)

    def test_full_compile(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            out = td / "out"
            self.make_source(source)
            result = compile_xlsform(source, out)
            self.assertTrue(result.output_xlsx.exists())
            release = load_workbook(result.output_xlsx, data_only=False)
            self.assertNotIn("methodmesh", release.sheetnames)
            ws = release["survey"]
            headers = [c.value for c in ws[1]]
            self.assertNotIn("mm_commit", headers)
            names = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
            self.assertIn("mm_authenticate_operator", names)
            self.assertIn("mm_create_attestation", names)
            self.assertIn("mm_submission_guard", names)

            intent_col = headers.index("body::intent") + 1
            auth_row = names.index("mm_authenticate_operator") + 2
            att_row = names.index("mm_create_attestation") + 2
            auth_intent = str(ws.cell(auth_row, intent_col).value)
            att_intent = str(ws.cell(att_row, intent_col).value)
            for intent in [auth_intent, att_intent]:
                self.assertIn("study_id=${mm_study_id}", intent)
                self.assertIn("form_id=${mm_form_id}", intent)
                self.assertIn("form_version=${mm_form_version}", intent)
                self.assertIn("form_instance_id=${mm_workflow_instance_id}", intent)
                self.assertNotIn("visit_id=${mm_workflow_instance_id}", intent)
            self.assertIn("operator_id=${mm_auth_credential_subject_id}", att_intent)
            self.assertIn("input_verification_execution_id=${mm_auth_methodmesh_execution_id}", att_intent)
            participant_row = names.index("participant_id") + 2
            relevant_col = headers.index("relevant") + 1
            readonly_col = headers.index("readonly") + 1
            self.assertIn("mm_auth_ok", ws.cell(participant_row, relevant_col).value)
            self.assertIn("mm_finalize_for_attestation", ws.cell(participant_row, readonly_col).value)
            age_row = names.index("age") + 2
            self.assertNotIn("no", ws.cell(age_row, readonly_col).value.lower())
            manifest = json.loads(result.manifest_json.read_text())
            self.assertEqual([m["name"] for m in manifest["commitment"]["members"]], ["participant_id", "age", "outcome"])
            explicit = [x["name"] for x in manifest["exclusions"] if x["explicit"]]
            self.assertEqual(explicit, ["ui_helper"])

            # ODK Validate requires every survey element to have a label or hint.
            # Generated hidden-answer return fields therefore receive a harmless hint.
            hint_cols = [i + 1 for i, h in enumerate(headers) if h == "hint" or (isinstance(h, str) and h.startswith("hint::"))]
            label_cols = [i + 1 for i, h in enumerate(headers) if h == "label" or (isinstance(h, str) and h.startswith("label::"))]
            appearance_col = headers.index("appearance") + 1
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, appearance_col).value == "hidden-answer":
                    has_label = any(str(ws.cell(r, c).value or "").strip() for c in label_cols)
                    has_hint = any(str(ws.cell(r, c).value or "").strip() for c in hint_cols)
                    self.assertTrue(has_label or has_hint, f"row {r} hidden-answer lacks label/hint")


    def test_release_version_is_generated_once_and_bound_everywhere(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            out = td / "out"
            self.make_source(source)
            fixed = datetime(2026, 9, 22, 18, 7, tzinfo=timezone(timedelta(hours=1)))
            with patch("methodmesh_xlsform.compiler._now_local", return_value=fixed):
                result = compile_xlsform(source, out)

            release = load_workbook(result.output_xlsx, data_only=False)
            settings = release["settings"]
            settings_headers = {str(c.value): c.column for c in settings[1] if c.value}
            self.assertEqual(settings.cell(2, settings_headers["version"]).value, "20260922180700")

            ws = release["survey"]
            headers = [c.value for c in ws[1]]
            names = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
            version_row = names.index("mm_form_version") + 2
            calc_col = headers.index("calculation") + 1
            self.assertEqual(ws.cell(version_row, calc_col).value, "'20260922180700'")

            manifest = json.loads(result.manifest_json.read_text())
            self.assertEqual(manifest["form"]["version"], "20260922180700")
            self.assertEqual(manifest["source"]["authoring_version"], "2026092201")
            self.assertEqual(manifest["build"]["compiled_at_local"], "2026-09-22T18:07:00+01:00")

    def test_source_version_may_be_absent_because_compiler_owns_release_version(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            out = td / "out"
            self.make_source(source)
            wb = load_workbook(source)
            settings = wb["settings"]
            headers = {str(c.value): c.column for c in settings[1] if c.value}
            settings.cell(2, headers["version"]).value = None
            wb.save(source)

            fixed = datetime(2026, 9, 22, 18, 8, tzinfo=timezone(timedelta(hours=1)))
            with patch("methodmesh_xlsform.compiler._now_local", return_value=fixed):
                result = compile_xlsform(source, out)
            release = load_workbook(result.output_xlsx, data_only=False)
            settings = release["settings"]
            settings_headers = {str(c.value): c.column for c in settings[1] if c.value}
            self.assertEqual(settings.cell(2, settings_headers["version"]).value, "20260922180800")
            manifest = json.loads(result.manifest_json.read_text())
            self.assertIsNone(manifest["source"]["authoring_version"])

    def test_release_bundle_contains_source_manifests_and_checksums(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            out = td / "out"
            self.make_source(source)
            fixed = datetime(2026, 9, 22, 18, 15, 42, tzinfo=timezone(timedelta(hours=1)))
            with patch("methodmesh_xlsform.compiler._now_local", return_value=fixed):
                result = compile_xlsform(source, out)

            self.assertTrue(result.release_dir.is_dir())
            self.assertEqual(result.release_dir.parent, out)
            self.assertIn("test_form__v20260922181542__build_20260922T181542", result.release_dir.name)
            self.assertEqual(result.source_copy.read_bytes(), source.read_bytes())
            for path in [result.output_xlsx, result.source_copy, result.manifest_json, result.human_manifest_md, result.report_md, result.checksums_sha256]:
                self.assertTrue(path.exists(), str(path))
                self.assertEqual(path.parent, result.release_dir)

            human = result.human_manifest_md.read_text()
            for heading in [
                "## Release identity",
                "## Source provenance",
                "## What the compiler added",
                "## Commitment policy",
                "## How to verify this release bundle",
                "## How to verify an ODK submission",
                "## Release discipline",
            ]:
                self.assertIn(heading, human)
            self.assertIn("Good provenance starts at compilation", human)

            manifest = json.loads(result.manifest_json.read_text())
            self.assertEqual(manifest["schema"], "methodmesh.xlsform_commitment_manifest.v2")
            self.assertEqual(manifest["bundle_schema"], "methodmesh.xlsform_release_bundle.v1")
            self.assertEqual(manifest["bundle"]["build_id"], result.release_dir.name)
            self.assertEqual(manifest["source"]["preserved_filename"], result.source_copy.name)
            self.assertEqual(len(manifest["compiler"]["implementation_sha256"]), 64)
            self.assertEqual(len(manifest["commitment"]["recipe_sha256"]), 64)

            checks = {}
            for line in result.checksums_sha256.read_text().splitlines():
                digest, filename = line.split("  ", 1)
                checks[filename] = digest
            self.assertEqual(set(checks), {
                result.source_copy.name, result.output_xlsx.name, result.manifest_json.name,
                result.human_manifest_md.name, result.report_md.name,
            })
            import hashlib
            for filename, expected in checks.items():
                actual = hashlib.sha256((result.release_dir / filename).read_bytes()).hexdigest()
                self.assertEqual(actual, expected)

    def test_repeated_builds_never_overwrite_previous_release_bundle(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            out = td / "out"
            self.make_source(source)
            fixed = datetime(2026, 9, 22, 18, 16, 5, tzinfo=timezone(timedelta(hours=1)))
            with patch("methodmesh_xlsform.compiler._now_local", return_value=fixed):
                first = compile_xlsform(source, out)
                second = compile_xlsform(source, out)
            self.assertNotEqual(first.release_dir, second.release_dir)
            self.assertTrue(first.release_dir.exists())
            self.assertTrue(second.release_dir.exists())
            self.assertTrue(second.release_dir.name.endswith("__02"))
            self.assertEqual(first.source_copy.read_bytes(), second.source_copy.read_bytes())

    def test_sentinel_form_compile_task_emits_machine_receipt_and_covers_it_with_checksums(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            out = td / "out"
            self.make_source(source)
            fixed = datetime(2026, 9, 22, 18, 20, 7, tzinfo=timezone(timedelta(hours=1)))
            with patch("methodmesh_xlsform.compiler._now_local", return_value=fixed), \
                 patch("methodmesh_xlsform.sentinel._now_local", side_effect=[fixed, fixed]):
                task = run_form_compile_task(source, out, request_id="req-test-001", request_source="test")

            self.assertEqual(task.status, "succeeded")
            self.assertEqual(task.task_type, "form.compile")
            self.assertTrue(task.event_json.exists())
            event = json.loads(task.event_json.read_text())
            self.assertEqual(event["schema"], "methodmesh.sentinel.task_event.v1")
            self.assertEqual(event["task_schema_version"], "1")
            self.assertEqual(event["request"]["request_id"], "req-test-001")
            self.assertEqual(event["request"]["input"]["filename"], source.name)
            self.assertEqual(len(event["request"]["input"]["sha256"]), 64)
            self.assertEqual(event["execution"]["status"], "succeeded")
            self.assertEqual(event["execution"]["implementation"]["version"], "0.4.0")
            self.assertEqual(event["result"]["form"]["form_version"], "20260922182007")

            manifest = json.loads(task.compile_result.manifest_json.read_text())
            self.assertEqual(manifest["sentinel_task"]["request_id"], "req-test-001")
            self.assertEqual(manifest["bundle"]["contents"]["sentinel_task_event"], "sentinel_task_event.json")

            checks = {}
            for line in task.compile_result.checksums_sha256.read_text().splitlines():
                digest, filename = line.split("  ", 1)
                checks[filename] = digest
            self.assertIn("sentinel_task_event.json", checks)
            import hashlib
            self.assertEqual(
                checks["sentinel_task_event.json"],
                hashlib.sha256(task.event_json.read_bytes()).hexdigest(),
            )

    def test_failed_sentinel_task_emits_failure_receipt(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            out = td / "out"
            missing = td / "missing.xlsx"
            with self.assertRaises(ValidationError):
                run_form_compile_task(missing, out, request_id="req-fail-001", request_source="test")
            event_path = out / "_sentinel_failed_tasks" / "req-fail-001.json"
            self.assertTrue(event_path.exists())
            event = json.loads(event_path.read_text())
            self.assertEqual(event["execution"]["status"], "failed")
            self.assertEqual(event["result"]["error_type"], "ValidationError")
            self.assertEqual(event["request"]["input"]["sha256"], None)

    def test_evidence_first_authentication_contract(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            out = td / "out"
            self.make_source(source)
            result = compile_xlsform(source, out)

            release = load_workbook(result.output_xlsx, data_only=False)
            ws = release["survey"]
            headers = [c.value for c in ws[1]]
            names = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]

            self.assertNotIn("mm_trusted_issuer_key_ids", names)
            for name in [
                "mm_auth_credential_id",
                "mm_auth_credential_subject_id",
                "mm_auth_issuer_key_id",
                "mm_auth_issuer_public_key_fingerprint_sha256",
                "mm_auth_issuer_public_key_base64",
                "mm_auth_issuer_signature_valid",
                "mm_auth_pin_verified",
                "mm_auth_tag_uid_hex",
                "mm_auth_verification_evidence_hash",
                "mm_auth_methodmesh_execution_id",
                "mm_auth_methodmesh_full_json",
            ]:
                self.assertIn(name, names)

            calculation_col = headers.index("calculation") + 1
            auth_ok_row = names.index("mm_auth_ok") + 2
            auth_calc = ws.cell(auth_ok_row, calculation_col).value
            self.assertIn("mm_auth_credential_verified", auth_calc)
            self.assertIn("mm_auth_pin_verified", auth_calc)
            self.assertIn("mm_auth_issuer_signature_valid", auth_calc)
            self.assertNotIn("issuer_trust_status", auth_calc)
            self.assertNotIn("trusted_issuer", auth_calc)

            intent_col = headers.index("body::intent") + 1
            auth_group_row = names.index("mm_authenticate_operator") + 2
            auth_intent = ws.cell(auth_group_row, intent_col).value
            self.assertIn("form_instance_id=${mm_workflow_instance_id}", auth_intent)
            self.assertNotIn("input_trusted_issuer_key_ids", auth_intent)

            canonical_row = names.index("mm_canonical_commitment_live") + 2
            canonical_calc = ws.cell(canonical_row, calculation_col).value
            self.assertIn("issuer_public_key_fingerprint_sha256", canonical_calc)
            self.assertIn("mm_auth_verification_evidence_hash", canonical_calc)

            manifest = json.loads(result.manifest_json.read_text())
            mm = manifest["methodmesh"]
            self.assertEqual(mm["field_authentication_mode"], "cryptographic_credential_evidence")
            self.assertEqual(mm["study_authorisation_mode"], "central_reconciliation")
            self.assertFalse(mm["issuer_allow_list_required_on_field_device"])
            self.assertEqual(
                mm["attestation_bound_authentication_evidence"],
                [
                    "credential_id_sha256",
                    "credential_subject_id_sha256",
                    "issuer_public_key_fingerprint_sha256",
                    "verification_evidence_hash",
                ],
            )

    def test_generated_return_fields_have_odk_label_or_hint(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            out = td / "out"
            self.make_source(source)
            result = compile_xlsform(source, out)
            release = load_workbook(result.output_xlsx, data_only=False)
            ws = release["survey"]
            headers = {c.value: c.column for c in ws[1] if c.value}
            for r in range(2, ws.max_row + 1):
                name = ws.cell(r, headers["name"]).value
                if not name or not (str(name).startswith("mm_auth_") or str(name).startswith("mm_att_")):
                    continue
                type_ = str(ws.cell(r, headers["type"]).value or "")
                if type_ not in {"text", "integer"}:
                    continue
                label = ws.cell(r, headers["label"]).value if "label" in headers else None
                hint = ws.cell(r, headers["hint"]).value if "hint" in headers else None
                self.assertTrue(str(label or "").strip() or str(hint or "").strip(), name)

    def test_repeat_requires_explicit_exclusion(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            source = td / "source.xlsx"
            self.make_source(source)
            wb = load_workbook(source)
            ws = wb["survey"]
            ws.insert_rows(2, 3)
            ws.cell(2, 1, "begin_repeat"); ws.cell(2, 2, "people")
            ws.cell(3, 1, "text"); ws.cell(3, 2, "person_name"); ws.cell(3, 3, "Name")
            ws.cell(4, 1, "end_repeat")
            wb.save(source)
            with self.assertRaises(ValidationError):
                compile_xlsform(source, td / "out")


if __name__ == "__main__":
    unittest.main()
